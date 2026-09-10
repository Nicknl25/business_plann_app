"""Turn 18 forward-proof readback (VS, 2026-09-10): fresh DB connection, the
live two-owner draft's stored roster, financials, planning run, run vitals,
and the workbook's Checks!B2. Read-only."""
import glob, json, os, sys, datetime
ROOT = r"C:\dev\business_plann_app"
sys.path.insert(0, r"C:/Users/IGNATI~1/AppData/Local/Temp/claude/C--dev-business-plann-app/8a0ab62f-278e-4bdb-9881-0fea4d8ebc67/scratchpad")
from pd_db import q

prefix = sys.argv[1]
print(f"=== readback {prefix} at {datetime.datetime.now().isoformat(timespec='seconds')}")
rows = q("SELECT draft_id, client_id, business_name, status, created_at, updated_at, people_json, financials_json FROM intake_consult_drafts WHERE draft_id LIKE %s", (prefix + "%",))
assert len(rows) == 1, rows
r = rows[0]
print("draft", r["draft_id"], r["business_name"], "client", r["client_id"], "status", r["status"], "created", r["created_at"], "updated", r["updated_at"])
ppl = json.loads(r["people_json"] or "{}") or {}
fin = json.loads(r["financials_json"] or "{}") or {}
people = [p for p in (ppl.get("people") or []) if isinstance(p, dict)]
print("people rows:", len(people))
for i, p in enumerate(people):
  print(f"  [{i}] full_name={p.get('full_name')!r} role_title={p.get('role_title')!r} annual_wage={p.get('annual_wage')!r} wage_source={p.get('wage_source')!r} experience_years={p.get('experience_years')!r} raw_keys={sorted(k for k in p if k in ('name','title','annual_pay','years_experience','education_credentials'))}")
print("rest_of_team_payroll_year1:", ppl.get("rest_of_team_payroll_year1"))
for k in ("current_payroll", "payroll_total_year1", "owner_compensation", "payroll_stated_total_target", "_owner_wage_conflict_hold", "_derived_patch_receipt", "current_revenue"):
  print(f"financials.{k} =", fin.get(k))
basis = fin.get("payroll_basis_people_roles")
print("payroll_basis_people_roles rows:", len(basis) if isinstance(basis, list) else basis)
if isinstance(basis, list):
  for b in basis:
    print("   ", {k: b.get(k) for k in ("full_name", "role_title", "annual_wage", "wage_source", "year1_payroll_amount") if k in b})
# the boundary contract on the live row
sys.path.insert(0, os.path.join(ROOT, "python")); sys.path.insert(0, os.path.join(ROOT, "python", "client_intake_and_finmo"))
from client_intake_and_finmo.post_intake_contracts.people_json_contract import PeopleJsonContract
from pydantic import ValidationError
try:
  PeopleJsonContract.model_validate(ppl); print("PeopleJsonContract on the live people_json: PASSED")
except ValidationError as e:
  print("PeopleJsonContract on the live people_json: errors", len(e.errors()), [er["loc"] for er in e.errors()][:6])
for pr in q("SELECT planning_run_id, run_status, current_stage, current_stage_status, failure_reason, started_at, completed_at FROM planning_runs WHERE draft_id LIKE %s ORDER BY started_at", (prefix + "%",)):
  print("planning_run", pr["planning_run_id"], pr["run_status"], pr["current_stage"], pr["current_stage_status"], pr["started_at"], "->", pr["completed_at"], "|", str(pr["failure_reason"] or "")[:300])
try:
  for v in q("SELECT * FROM run_vitals_runs WHERE draft_id LIKE %s ORDER BY id DESC LIMIT 2", (prefix + "%",)):
    print("run_vitals_runs", {k: (str(val)[:60] if val is not None else None) for k, val in v.items()})
except Exception as e:
  print("run_vitals_runs query failed:", e)
# workbooks written for this business today
name = str(r["business_name"] or "")
stem = name.replace("&", "").replace("  ", " ")
for d in (r"C:\dev\Cilient Plans", r"C:\Users\IgnatiusHenry\OneDrive - Tithe Financial Wealth Management\Apps\Business Plan Generator Sources\Client Plans\Financial Models"):
  hits = sorted(glob.glob(os.path.join(d, f"{name.split(' ')[0]}*09-10-2026*.xlsx")))
  print("workbooks in", d, ":", [os.path.basename(h) for h in hits])
  for h in hits:
    try:
      from openpyxl import load_workbook
      wb = load_workbook(h, data_only=True, read_only=True)
      if "Checks" in wb.sheetnames:
        ws = wb["Checks"]
        print("   Checks!B2 =", ws["B2"].value)
        fails = []
        for row in ws.iter_rows(min_row=1, max_row=200, values_only=True):
          if row and any(isinstance(c, str) and c.strip().upper() == "FAIL" for c in row):
            fails.append(row[:4])
        print("   FAIL rows:", fails)
      else:
        print("   no Checks sheet; sheets:", wb.sheetnames[:12])
      if "Valuation" in wb.sheetnames:
        ws = wb["Valuation"]
        for row in ws.iter_rows(min_row=1, max_row=60, values_only=True):
          if row and any(isinstance(c, str) and ("owner" in c.lower() or "sde" in c.lower() or "equity value" in c.lower()) for c in row):
            print("   Valuation:", [c for c in row[:8]])
      wb.close()
    except Exception as e:
      print("   workbook read failed:", e)
