"""Read the neighbor workbooks vs_r7_neighbors.py built and Excel-recalculated
(scratch dirs before_<id> / after_<id>): Checks!B2, every FAIL row on Checks,
the Valuation SDE row (note, Q1, Total), TV at the exit multiple, equity value,
implied multiple. No rebuild, no Excel - openpyxl data_only on the saved files."""
import glob, os, sys
import openpyxl

base = sys.argv[1]
for label in ("before", "after"):
  for prefix in ("46ae584a", "280a55e1", "3201a64c"):
    files = sorted(glob.glob(os.path.join(base, f"{label}_{prefix}", "*.xlsx")))
    if not files:
      print(label, prefix, "NO FILE"); continue
    path = files[-1]
    w = openpyxl.load_workbook(path, data_only=True)
    ws = w["Checks"]
    fails = []
    for r in range(7, ws.max_row + 1):
      vals = [ws.cell(row=r, column=c).value for c in range(1, 10)]
      if "FAIL" in vals:
        fails.append((r, vals[0], vals[1]))
    checks_total = sum(1 for r in range(7, ws.max_row + 1) if ws.cell(row=r, column=9).value in ("OK", "FAIL", "PASS"))
    v = w["Valuation"]
    rows = {}
    for r in range(1, v.max_row + 1):
      a = v.cell(row=r, column=1).value
      if isinstance(a, str):
        rows[a] = r
    def cell(label_prefix, col):
      for lab, r in rows.items():
        if lab.startswith(label_prefix):
          return v.cell(row=r, column=col).value
      return "LABEL NOT FOUND"
    print(f"=== {label} {prefix} {os.path.basename(path)}")
    print(f"    Checks!B2 = {ws['B2'].value!r}  (status rows counted: {checks_total}; FAIL rows: {fails if fails else 'none'})")
    print(f"    SDE note  = {cell(chr(83)+'eller', 2)!r}")
    print(f"    SDE Q1    = {cell(chr(83)+'eller', 3)!r}   SDE Total = {cell(chr(83)+'eller', 23)!r}")
    print(f"    TV exit multiple = {cell('Terminal value — exit multiple', 2)!r}")
    print(f"    Equity value     = {cell('Equity value', 2)!r}")
    print(f"    Implied multiple = {cell('Implied multiple', 2)!r}")
