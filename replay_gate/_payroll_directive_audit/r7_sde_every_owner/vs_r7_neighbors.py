"""Item 7 / R4 neighbor check (VS, 2026-09-10).

For each named neighbor draft: read the stored row (read-only), build the
workbook OFFLINE into a scratch directory (never Cilient Plans), recalculate
it with Excel COM in a SEPARATE subprocess (the valuation guard's recipe;
in-process COM after the MySQL connector + openpyxl build died with an
access violation on this machine, the standalone recipe does not), and read
  * Checks!B2 (the model status - the standing rule),
  * the Valuation sheet's SDE row note, Q1 formula/value and Total,
  * the terminal value at the exit multiple and the equity value,
  * the fact twin's owner_comp_q / sde_y5 / equity_value
    (writing_phase.facts.valuation), and the stored mirror.

Usage: python -u vs_r7_neighbors.py --root <repo-or-worktree> --label before|after --out <dir>
The --root decides which code builds the workbook, so the same script
records HEAD (before) and the edited tree (after).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import subprocess
import sys

DRAFTS = [
  ("46ae584a", "Bellweather Auto Works (standing rule)"),
  ("280a55e1", "Sunny_V3 canary 2026-09-09 22:20"),
  ("3201a64c", "Marchetti & Fen LIVE two-owner (read-only)"),
]

RECALC_SRC = r'''
import sys, time
import pythoncom, win32com.client as win32
path = sys.argv[1]
pythoncom.CoInitialize()
x = win32.Dispatch("Excel.Application")
x.Visible = False; x.DisplayAlerts = False
wb = x.Workbooks.Open(path)
for _ in range(20):
  try:
    wb.Sheets(1).Name; break
  except Exception:
    time.sleep(1.5)
x.CalculateFullRebuild()
wb.Save(); wb.Close(False); x.Quit()
pythoncom.CoUninitialize()
print("RECALC_OK")
'''


def _recalc(path: str) -> str:
  if os.environ.get("R7_SKIP_RECALC"):
    return "skipped by R7_SKIP_RECALC"
  helper = os.path.join(os.path.dirname(path), "_recalc_helper.py")
  with open(helper, "w", encoding="utf-8") as fh:
    fh.write(RECALC_SRC)
  p = subprocess.run([sys.executable, "-u", helper, path], capture_output=True, text=True, timeout=600)
  if p.returncode == 0 and "RECALC_OK" in (p.stdout or ""):
    return "excel_com_full_rebuild (subprocess)"
  tail = ((p.stderr or p.stdout or "").strip().splitlines() or ["no output"])[-1]
  return f"FAILED rc={p.returncode}: {tail[:200]}"


def main() -> int:
  print("vs_r7_neighbors start", flush=True)
  ap = argparse.ArgumentParser()
  ap.add_argument("--root", required=True)
  ap.add_argument("--label", required=True)
  ap.add_argument("--out", required=True)
  a = ap.parse_args()
  root = os.path.abspath(a.root)
  sys.path[:0] = [root, os.path.join(root, "python"), os.path.join(root, "python", "client_intake_and_finmo")]
  from dotenv import load_dotenv
  load_dotenv(os.path.join("C:/dev/business_plann_app", ".env"))
  from intake_submission import get_mysql_connection  # type: ignore
  from client_statements_output_excel.export_client_workbook import export_workbook_for_row  # type: ignore
  import client_statements_output_excel as _pkg  # type: ignore
  from writing_phase.facts import valuation as V  # type: ignore
  import openpyxl

  got = os.path.normcase(os.path.abspath(os.path.dirname(_pkg.__file__)))
  want = os.path.normcase(os.path.join(root, "client_statements_output_excel"))
  print(f"== neighbor check label={a.label} root={root}", flush=True)
  print(f"   workbook package resolved from {got} (provenance {'OK' if got == want else 'FAIL'})", flush=True)
  if got != want:
    return 3
  os.makedirs(a.out, exist_ok=True)
  conn = get_mysql_connection()
  cur = conn.cursor(dictionary=True)
  for prefix, name in DRAFTS:
    cur.execute("SELECT * FROM intake_consult_drafts WHERE draft_id LIKE %s", (prefix + "%",))
    row = cur.fetchone()
    if not row:
      print(f"{prefix} {name}: NOT FOUND", flush=True); continue
    fin = json.loads(row.get("financials_json") or "{}")
    ppl = json.loads(row.get("people_json") or "{}")
    owners = [(p.get("full_name"), p.get("role_title"), p.get("annual_wage")) for p in ppl.get("people") or []]
    out_dir = os.path.join(a.out, f"{a.label}_{prefix}")
    os.makedirs(out_dir, exist_ok=True)
    path = str(export_workbook_for_row(dict(row), output_dir=out_dir))
    built_digest = hashlib.sha256(open(path, "rb").read()).hexdigest()[:12]
    recalc = _recalc(path)
    w = openpyxl.load_workbook(path, data_only=True)
    checks_b2 = w["Checks"]["B2"].value if "Checks" in w.sheetnames else "NO CHECKS SHEET"
    sde = {}
    if "Valuation" in w.sheetnames:
      ws = w["Valuation"]
      for r in range(1, ws.max_row + 1):
        lab = ws.cell(row=r, column=1).value
        if isinstance(lab, str) and lab.startswith("Seller's discretionary earnings"):
          sde = {"note": ws.cell(row=r, column=2).value, "q1": ws.cell(row=r, column=3).value,
                 "total": ws.cell(row=r, column=23).value}
        if isinstance(lab, str) and lab.startswith("Equity value"):
          sde["equity_value"] = ws.cell(row=r, column=2).value
        if isinstance(lab, str) and lab.startswith("Terminal value") and "exit multiple" in lab:
          sde["tv_exit_multiple"] = ws.cell(row=r, column=2).value
        if isinstance(lab, str) and lab.startswith("Implied multiple"):
          sde["implied_multiple"] = ws.cell(row=r, column=2).value
    wf = openpyxl.load_workbook(path, data_only=False)
    sde_formula_q1 = None
    if "Valuation" in wf.sheetnames:
      ws = wf["Valuation"]
      for r in range(1, ws.max_row + 1):
        lab = ws.cell(row=r, column=1).value
        if isinstance(lab, str) and lab.startswith("Seller's discretionary earnings"):
          sde_formula_q1 = ws.cell(row=r, column=3).value
    fact = V.compute_valuation(conn.cursor(), dict(row))
    print(f"--- {prefix} {name}", flush=True)
    print(f"    owners_in_people_rows: {owners}", flush=True)
    print(f"    stored mirror financials.owner_compensation (monthly): {fin.get('owner_compensation')}", flush=True)
    print(f"    workbook: {os.path.basename(path)} built sha256[:12]={built_digest} recalc={recalc} bytes_after={os.path.getsize(path)}", flush=True)
    print(f"    sheets = {w.sheetnames}", flush=True)
    print(f"    Checks!B2 = {checks_b2!r}", flush=True)
    print(f"    Valuation SDE note = {sde.get('note')!r}", flush=True)
    print(f"    Valuation SDE Q1 formula = {sde_formula_q1!r}", flush=True)
    print(f"    Valuation SDE Q1 value = {sde.get('q1')!r}  Total = {sde.get('total')!r}", flush=True)
    print(f"    Valuation TV(exit multiple) = {sde.get('tv_exit_multiple')!r}  Equity value = {sde.get('equity_value')!r}  Implied multiple = {sde.get('implied_multiple')!r}", flush=True)
    print(f"    fact twin owner_comp_q = {fact.get('owner_comp_q')!r}  sde_y5 = {fact.get('sde_y5')!r}  equity_value = {fact.get('equity_value')!r}", flush=True)
  conn.close()
  return 0


if __name__ == "__main__":
  raise SystemExit(main())
