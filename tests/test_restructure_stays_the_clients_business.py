"""The restructure stays the client's business (Nick 2026-09-11).

"A client can be handed a plan for a business they didn't describe, with three
lines they never mentioned, and nothing tells them. That's not a restructure,
that's a different company." Five rules:
  1. new lines only if the client named them;
  2. no dropping a client's line without asking;
  3. price and volume ceilings tied to market data;
  4. a failed review blocks the design, never approves it;
  5. a restructured plan is always labelled as one.
"""
from __future__ import annotations

import inspect
import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
for p in (ROOT, os.path.join(ROOT, "python")):
  if p not in sys.path:
    sys.path.insert(0, p)
_EXTRA = os.path.join(ROOT, "python", "client_intake_and_finmo")
if _EXTRA not in sys.path:
  sys.path.append(_EXTRA)

import api_handlers.intake_consult as IC  # noqa: E402
from client_intake_and_finmo.post_intake_restructure import client_bounds as CB  # noqa: E402
from client_intake_and_finmo.post_intake_restructure import constraint_author as CA  # noqa: E402

CLIENT = [
  "We sell loose-leaf tea blends online and at two farmers markets.",
  "At some point we'd like to start a wholesale program for local cafes, maybe next year.",
]
BOUNDS = {
  "feasible_region_exists": True,
  "existing_lines": [
    {"lob": "Retail", "product": "Loose-leaf blends", "price_multiplier_max": 2.4,
     "volume_multiplier_max": 3.0, "can_drop": False, "rationale": "core"},
    {"lob": "Markets", "product": "Farmers market stall", "price_multiplier_max": 1.3,
     "volume_multiplier_max": 1.5, "can_drop": True, "rationale": "low margin"},
  ],
  "new_line_candidates": [
    {"lob": "Wholesale", "product": "Cafe wholesale program", "unit_price": 18.0,
     "q11_quarterly_revenue_max": 40000.0,
     "client_quote": "we'd like to start a wholesale program for local cafes"},
    {"lob": "Subscriptions", "product": "Monthly tea subscription box", "unit_price": 30.0,
     "q11_quarterly_revenue_max": 60000.0,
     "client_quote": "customers would love a monthly subscription box"},
    {"lob": "Events", "product": "Tea tasting events", "unit_price": 45.0,
     "q11_quarterly_revenue_max": 15000.0, "client_quote": ""},
  ],
}


def _bound(market=None):
  return CB.bound_to_the_clients_business(BOUNDS, client_statements=CLIENT, market_ceiling=market)


class RuleOneNewLinesOnlyIfTheClientNamedThem(unittest.TestCase):
  def test_a_real_quote_naming_the_line_verifies(self):
    ok, why = CB.quote_is_the_clients(
      "we'd like to start a wholesale program for local cafes", CLIENT, "Wholesale Cafe wholesale program")
    self.assertTrue(ok, why)

  def test_words_the_client_never_typed_are_not_a_quote(self):
    ok, why = CB.quote_is_the_clients(
      "customers would love a monthly subscription box", CLIENT, "Monthly tea subscription box")
    self.assertFalse(ok)
    self.assertEqual(why, "quote_not_in_the_clients_words")

  def test_no_quote_no_line(self):
    self.assertEqual(CB.quote_is_the_clients("", CLIENT, "Tea tasting events")[1], "no_client_quote")

  def test_a_real_quote_about_something_else_does_not_name_the_line(self):
    ok, why = CB.quote_is_the_clients(
      "We sell loose-leaf tea blends online", CLIENT, "Cafe wholesale program")
    self.assertFalse(ok)
    self.assertEqual(why, "quote_does_not_name_the_line")

  def test_unnamed_lines_never_reach_the_search(self):
    b = _bound(market=lambda entry, kind: 20000.0)
    names = [nl["product"] for nl in b["new_line_candidates"]]
    self.assertEqual(names, ["Cafe wholesale program"])
    dropped = {d["name"]: d["reason"] for d in b["client_bounds"]["new_lines_dropped"]}
    self.assertEqual(dropped["Subscriptions Monthly tea subscription box"], "quote_not_in_the_clients_words")
    self.assertEqual(dropped["Events Tea tasting events"], "no_client_quote")


class RuleTwoNoDroppingWithoutAsking(unittest.TestCase):
  def test_no_line_can_be_dropped_and_the_proposal_is_kept_as_a_question(self):
    b = _bound()
    self.assertFalse(any(l["can_drop"] for l in b["existing_lines"]))
    self.assertEqual([d["line"] for d in b["client_bounds"]["drops_proposed"]],
                     ["Markets/Farmers market stall"])


class RuleThreeCeilingsTiedToMarketData(unittest.TestCase):
  def test_without_market_data_there_is_no_headroom(self):
    b = _bound()
    for line in b["existing_lines"]:
      self.assertEqual(line["price_multiplier_max"], 1.0)
      self.assertEqual(line["volume_multiplier_max"], 1.0)
    self.assertTrue(all(c["source"] == "no_market_data" for c in b["client_bounds"]["price_capped"]))

  def test_without_market_data_even_a_client_named_line_is_not_sized(self):
    b = _bound()
    self.assertEqual(b["new_line_candidates"], [])
    reasons = {d["reason"] for d in b["client_bounds"]["new_lines_dropped"]}
    self.assertIn("no_market_data", reasons)

  def test_a_market_figure_caps_the_judgment_and_never_raises_it(self):
    b = _bound(market=lambda entry, kind: {"price": 1.2, "volume": 5.0}.get(kind, 20000.0))
    core = b["existing_lines"][0]
    self.assertEqual(core["price_multiplier_max"], 1.2)   # market 1.2 < judged 2.4
    self.assertEqual(core["volume_multiplier_max"], 3.0)  # judged 3.0 < market 5.0
    self.assertEqual(b["new_line_candidates"][0]["q11_quarterly_revenue_max"], 20000.0)

  def test_applying_twice_changes_nothing_and_keeps_the_record(self):
    once = _bound()
    twice = CB.bound_to_the_clients_business(once, client_statements=CLIENT)
    self.assertEqual(once, twice)


class RuleFourAFailedReviewBlocks(unittest.TestCase):
  def test_no_review_is_not_an_approval(self):
    src = inspect.getsource(IC.post_intake_consult_system_run_handler)
    self.assertIn("if _rs_review is None:", src)
    self.assertIn('_rs_blocked_reason = "review_unavailable"', src)
    self.assertNotIn("if _rs_review is not None and not bool(_rs_review.get(\"approved\")):", src)


class RuleFiveAlwaysLabelled(unittest.TestCase):
  SRC = inspect.getsource(IC.post_intake_consult_system_run_handler)

  def test_the_sentence_is_gone(self):
    self.assertNotIn("simply IS the forecast", self.SRC)

  def test_the_workbook_the_email_and_the_writing_trigger_all_read_the_mark(self):
    self.assertIn("_RS_LABEL", self.SRC)
    self.assertIn("RESTRUCTURED PLAN, NOT THE BUSINESS AS DESCRIBED", self.SRC)
    self.assertIn("WRITING_PHASE_WITHHELD", self.SRC)
    guard = self.SRC.find("if _rs_delivered_restructure:\n        # ALWAYS LABELLED (Nick 2026-09-11): the written plan")
    trigger = self.SRC.find("_auto_trigger_writing_phase(app, diagnostic_payload, result_draft_id)")
    self.assertGreater(guard, 0)
    self.assertGreater(trigger, guard)

  # The Vespertine shape (live run 2026-09-11): the design asked for a team
  # payroll of 360,000, the real model kept the client's payroll (named
  # people's pay is never cut), and rent fell 16,200 -> 13,500 a quarter.
  DESIGN = {"team": {"annual_payroll": 360000.0}, "facility": {"quarterly_rent_target": 13500.0}}

  @staticmethod
  def _finmo(rent_q, cogs_q):
    rows = [{"quarter_index": 0}]
    for q in range(1, 21):
      rows.append({"quarter_index": q, "revenue": 100000.0, "payroll": 137250.0,
                   "lease_rent": rent_q, "cogs": cogs_q, "marketing": 6000.0,
                   "g_and_a": 8000.0, "ebitda": -60000.0})
    return {"quarter_rows": rows}

  @staticmethod
  def _mi(price, new_line=False):
    rows = [{"lob": "Retail", "product": "Loose-leaf blends", "driver": "Unit Price", "values": [price] * 21},
            {"lob": "Retail", "product": "Loose-leaf blends", "driver": "Capacity", "values": [1000.0] * 21}]
    if new_line:
      rows.append({"lob": "Wholesale", "product": "Cafe wholesale program", "driver": "Unit Price",
                   "values": [18.0] * 21})
    return {"sections": {"revenue": rows}}

  def _rows(self):
    b = _bound(market=lambda entry, kind: {"price": (1.2, CB.INTAKE_PRICE_SOURCE)}.get(kind, 20000.0))
    before = {"finmo": self._finmo(16200.0, 37900.0), "model_input": self._mi(20.0)}
    after = {"finmo": self._finmo(13500.0, 34600.0), "model_input": self._mi(22.0, new_line=True)}
    return CB.restructure_changes(self.DESIGN, b["client_bounds"], before=before, after=after)

  def test_a_design_target_that_never_landed_is_never_claimed(self):
    rows = self._rows()
    self.assertFalse(any("360,000" in (r["restructured"] + r["as_described"]) for r in rows))
    payroll = {r["area"]: r for r in rows}["Payroll cost, incl. employer costs (year 1)"]
    self.assertEqual((payroll["as_described"], payroll["restructured"], payroll["note"]),
                     ("$549,000", "$549,000", "unchanged"))

  def test_every_change_is_a_row_with_both_sides_from_the_models(self):
    rows = {r["area"]: r for r in self._rows()}
    rent = rows["Rent (year 3)"]
    self.assertEqual((rent["as_described"], rent["restructured"]), ("$64,800", "$54,000"))
    self.assertEqual(rows["Cost of goods (year 1)"]["note"], "37.9% of revenue -> 34.6%")
    price = rows["Loose-leaf blends: price (year 3)"]
    self.assertEqual((price["as_described"], price["restructured"]), ("$20", "$22"))
    self.assertIn("price ceiling judged at intake", price["note"])
    self.assertIn("wholesale program for local cafes", rows["New line: Cafe wholesale program"]["note"])
    self.assertIn("NOT applied", rows["Markets/Farmers market stall"]["note"])

  def test_the_email_tells_the_same_story_as_the_sheet(self):
    rows = self._rows()
    text = CB.restructure_label_text(rows)
    self.assertIn("NOT THE BUSINESS AS DESCRIBED", text)
    for r in rows:
      self.assertIn(r["area"], text)
    self.assertIn("What Changed sheet", text)


class TheIntakePriceBandIsTheMarketFigureTests(unittest.TestCase):
  """Nick 2026-09-11: use the intake's GPT-estimated price bands as the market
  figure for now. The fact shape is Ferriday & Blythe's (73a71cfe), as stored."""
  FIN = {"_coherence": {"price_market_facts": {
    "Primary line of business␟Surgical procedures": {"ceiling_dollars": 2590.0, "market_slice_hash": "x"},
    "Primary line of business␟Boarding stays": {"ceiling_dollars": 69.6, "market_slice_hash": "x"},
  }}}
  PRICES = {"Primary line of business/Surgical procedures": 2072.0,
            "Primary line of business/Boarding stays": 72.0}

  def _line(self, product, pmax=2.0):
    return {"lob": "Primary line of business", "product": product,
            "price_multiplier_max": pmax, "volume_multiplier_max": 1.5, "can_drop": False}

  def test_the_ceiling_is_the_band_over_the_current_price(self):
    ceiling = CB.intake_price_ceiling(self.FIN, self.PRICES)
    value, source = ceiling(self._line("Surgical procedures"), "price")
    self.assertAlmostEqual(value, 2590.0 / 2072.0)
    self.assertEqual(source, CB.INTAKE_PRICE_SOURCE)

  def test_the_band_gives_pricing_headroom_the_judgment_cannot_exceed(self):
    b = CB.bound_to_the_clients_business(
      {"existing_lines": [self._line("Surgical procedures", 2.0), self._line("Boarding stays", 1.3)]},
      client_statements=[], market_ceiling=CB.intake_price_ceiling(self.FIN, self.PRICES))
    surgery, boarding = b["existing_lines"]
    self.assertAlmostEqual(surgery["price_multiplier_max"], 2590.0 / 2072.0)   # the band, not GPT's 2.0
    self.assertEqual(boarding["price_multiplier_max"], 1.0)   # already above its band: no raise
    self.assertEqual(surgery["volume_multiplier_max"], 1.0)   # no intake figure for volume
    self.assertEqual([h["source"] for h in b["client_bounds"]["headroom_allowed"]], [CB.INTAKE_PRICE_SOURCE])

  def test_no_intake_band_means_no_market_figure(self):
    self.assertIsNone(CB.intake_price_ceiling({}, self.PRICES))
    self.assertIsNone(CB.intake_price_ceiling(self.FIN, self.PRICES)(self._line("Surgical procedures"), "volume"))

  def test_the_handler_feeds_the_band_to_every_bounding(self):
    src = inspect.getsource(IC.post_intake_consult_system_run_handler)
    self.assertIn("_rs_intake_price_ceiling(_rs_fin", src)
    self.assertEqual(src.count("market_ceiling=_rs_market"), 2)


class WhatChangedSheetTests(unittest.TestCase):
  def test_the_sheet_sits_after_the_cover_and_carries_every_row(self):
    from openpyxl import Workbook
    from client_statements_output_excel.what_changed_sheet import (
      WHAT_CHANGED_SHEET, build_what_changed_sheet)
    from client_statements_output_excel.cover_sheet import COVER_SHEET
    wb = Workbook()
    wb.active.title = COVER_SHEET
    wb.create_sheet("FINMO")
    rows = [{"area": "Rent (year 3)", "as_described": "$64,800",
             "restructured": "$54,000", "note": ""}]
    ws = build_what_changed_sheet(wb, rows)
    self.assertEqual(wb.sheetnames[:2], [COVER_SHEET, WHAT_CHANGED_SHEET])
    self.assertIn("NOT THE BUSINESS AS DESCRIBED", ws["A1"].value)
    self.assertEqual([ws.cell(row=5, column=c).value for c in range(1, 4)],
                     ["Rent (year 3)", "$64,800", "$54,000"])

  def test_only_a_restructured_workbook_gets_it(self):
    from client_statements_output_excel import export_client_workbook as E
    src = inspect.getsource(E.export_workbook_for_row)
    self.assertIn("if what_changed is not None:", src)
    handler = inspect.getsource(IC.post_intake_consult_system_run_handler)
    self.assertIn('what_changed=_rs_delivered_restructure.get("rows")', handler)


class TheIntakeWalkIsUntouchedTests(unittest.TestCase):
  def test_the_default_tool_has_no_quote_field(self):
    items = CA._submit_tool(False)["function"]["parameters"]["properties"]["new_line_candidates"]["items"]
    self.assertNotIn("client_quote", items["properties"])
    self.assertIs(CA._submit_tool(False), CA._SUBMIT_TOOL)

  def test_the_restructure_tool_requires_the_quote(self):
    items = CA._submit_tool(True)["function"]["parameters"]["properties"]["new_line_candidates"]["items"]
    self.assertIn("client_quote", items["required"])

  def test_the_clients_words_only_appear_when_given(self):
    kw = dict(compact={}, stated_facts={}, current_structure={})
    self.assertNotIn("CLIENT'S OWN WORDS", CA._build_user_prompt(**kw))
    self.assertIn("CLIENT'S OWN WORDS", CA._build_user_prompt(**kw, client_statements=CLIENT))

  def test_the_validator_carries_the_quote(self):
    v = CA.validate_restructure_bounds(bounds={
      "feasible_region_exists": True, "existing_lines": [],
      "new_line_candidates": [{"lob": "W", "product": "Cafe wholesale", "unit_price": 10,
                               "q11_quarterly_revenue_max": 100, "client_quote": "a quote"}],
      "team": {"min_annual_payroll": 1}, "reality_constraints": {
        "real_market": "x", "real_physics": "x", "still_this_business": "x", "lender_defensible": "x"},
    }, stated_owner_annual_wage=0.0)
    self.assertEqual(v["new_line_candidates"][0]["client_quote"], "a quote")

  def test_the_restructure_path_uses_the_client_bounds_after_every_bounds_change(self):
    src = inspect.getsource(IC.post_intake_consult_system_run_handler)
    self.assertIn("require_client_quote=True", src)
    self.assertGreaterEqual(src.count("_rs_bound_to_client("), 2)


if __name__ == "__main__":
  unittest.main()
