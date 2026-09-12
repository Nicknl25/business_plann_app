"""NO WORKBOOK SHIPS UNVERIFIED (Nick 2026-09-11).

"Excel wasn't unreachable. It said 'call was rejected by callee' - that's
BUSY, not ABSENT. The code reads a transient rejection as 'Excel isn't
installed' and skips the gate. That misreading is the whole defect. And it
matters far more than one workbook, because in production there IS no
Excel. PythonAnywhere is Linux. So the branch written for 'Excel might not
be here' becomes the normal path, and the model-status gate turns itself
off permanently the day we deploy. Until then: a rejection is retried, a
genuine absence FAILS the run, and no workbook ships unverified."

Halvorsen Tide 836c2ca2, 2026-09-11 20:36:29: -2147418111
RPC_E_CALL_REJECTED on the Excel COM recalc, logged as
"workbook_model_status_check_skipped: recalc unavailable", and the raw
openpyxl workbook - no cached values, Checks!B2 never read - shipped to the
folder and was emailed four milliseconds later. Two tests in this repo
ASSERTED that outcome: a workbook with Checks!B2 = "FAIL" was specified to
pass when the recalc reported an environment error. They are deleted; these
replace them.

THE LAW (writing_phase/checks.py:11, written after the CoInitialize bug of
2026-08-29): A CHECK THAT CANNOT RUN FAILS THE SECTION. IT NEVER PASSES BY
DEFAULT. This module enforces it on the workbook gate.

THE ENGINE IS THE ANSWER, NOT THE RETRY. Excel COM is one engine, Windows
only. LibreOffice headless is the other, and the one production has. The
delivered workbooks use twenty standard functions and no volatile ones, and
carry fullCalcOnLoad=1, so a headless convert recalculates them - verified
by inventory on the 09-11 deliverables.
"""
from __future__ import annotations

import os
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
for p in (ROOT, os.path.join(ROOT, "python")):
  if p not in sys.path:
    sys.path.insert(0, p)

from client_intake_and_finmo.fail_fast.common import PostIntakePreconditionFailed  # noqa: E402
from client_intake_and_finmo.post_intake_runtime_validation import (  # noqa: E402
  workbook_model_status as wms,
)

# THE REAL WRAPPER, captured at import. pytest imports every test module at
# collection - before any test runs - so this is always the genuine
# function even when an earlier file in the same process monkeypatches
# wms._recalc_workbook_via_excel_com and forgets to put it back (two files
# did exactly that on 2026-09-11, and the retry tests below died on a
# one-argument lambda that lived in someone else's file).
_REAL_EXCEL_WRAPPER = wms._recalc_workbook_via_excel_com


def _workbook_with_status(value, tmp: Path, name="wb.xlsx") -> Path:
  import openpyxl
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Checks"
  ws.cell(row=2, column=1, value="Model Status")
  ws.cell(row=2, column=2, value=value)
  out = tmp / name
  wb.save(str(out))
  return out


def _fake_soffice(tmp: Path, *, behaviour: str = "convert") -> list:
  """A stand-in for `soffice --headless --convert-to xlsx --outdir D F`:
  'convert' writes D/basename(F) (a real convert would recalculate on the
  way through; here the input already carries the value under test);
  'silent' exits 0 and writes nothing - the shape of a convert that
  produced no file; 'crash' exits 1."""
  script = tmp / "fake_soffice.py"
  script.write_text(textwrap.dedent(f"""
    import shutil, sys, os
    args = sys.argv[1:]
    outdir = args[args.index("--outdir") + 1]
    src = args[-1]
    mode = {behaviour!r}
    if mode == "crash":
        sys.exit(1)
    if mode == "convert":
        os.makedirs(outdir, exist_ok=True)
        shutil.copyfile(src, os.path.join(outdir, os.path.basename(src)))
    sys.exit(0)
  """), encoding="utf-8")
  return [sys.executable, str(script)]


class ErrorsAreClassifiedBeforeTheyAreActedOn(unittest.TestCase):
  def test_call_rejected_by_callee_is_transient_not_absent(self):
    """THE misreading: -2147418111 is Excel saying BUSY."""
    self.assertEqual(
      wms.classify_recalc_error(
        "excel_com_failure: com_error: (-2147418111, 'Call was rejected by callee.', None, None)"),
      "transient")

  def test_retry_later_is_transient(self):
    self.assertEqual(
      wms.classify_recalc_error("excel_com_failure: com_error: (-2147417846, 'The message filter indicated that the application is busy.', None, None)"),
      "transient")

  def test_missing_pywin32_is_absent(self):
    self.assertEqual(
      wms.classify_recalc_error("pywin32_unavailable: ImportError: No module named 'win32com'"),
      "absent")

  def test_excel_not_registered_is_absent(self):
    self.assertEqual(
      wms.classify_recalc_error("excel_com_failure: com_error: (-2147221005, 'Invalid class string', None, None)"),
      "absent")

  def test_a_real_com_failure_is_a_failure(self):
    """Anything that is neither busy nor missing is a genuine failure of
    the recalc - never retried into a pass, never read as absence."""
    self.assertEqual(
      wms.classify_recalc_error("excel_com_failure: com_error: (-2147352567, 'Exception occurred.', (0, 'Microsoft Excel', 'Save method of Workbook class failed', ...), None)"),
      "failure")


class ABusySignalIsRetried(unittest.TestCase):
  def setUp(self):
    self._orig = wms._recalc_workbook_via_excel_com_unlocked

  def tearDown(self):
    wms._recalc_workbook_via_excel_com_unlocked = self._orig

  def test_transient_then_success_returns_success(self):
    calls = []
    def fake(path):
      calls.append(path)
      if len(calls) < 3:
        return "excel_com_failure: com_error: (-2147418111, 'Call was rejected by callee.', None, None)"
      return None
    wms._recalc_workbook_via_excel_com_unlocked = fake
    err = _REAL_EXCEL_WRAPPER("x.xlsx", attempts=3, backoff_seconds=0.0)
    self.assertIsNone(err)
    self.assertEqual(len(calls), 3)

  def test_transient_every_time_is_still_an_error_after_the_attempts(self):
    wms._recalc_workbook_via_excel_com_unlocked = (
      lambda p: "excel_com_failure: com_error: (-2147418111, 'Call was rejected by callee.', None, None)")
    err = _REAL_EXCEL_WRAPPER("x.xlsx", attempts=3, backoff_seconds=0.0)
    self.assertIsNotNone(err)
    self.assertIn("2147418111", err)

  def test_absence_is_not_retried(self):
    calls = []
    def fake(path):
      calls.append(path)
      return "pywin32_unavailable: ImportError: No module named 'win32com'"
    wms._recalc_workbook_via_excel_com_unlocked = fake
    _REAL_EXCEL_WRAPPER("x.xlsx", attempts=3, backoff_seconds=0.0)
    self.assertEqual(len(calls), 1, "a missing engine is not busy - retrying it is just failing slower")


class LibreOfficeIsAnEngine(unittest.TestCase):
  def setUp(self):
    self._tmp = tempfile.TemporaryDirectory(prefix="wb_engine_")
    self.tmp = Path(self._tmp.name)

  def tearDown(self):
    self._tmp.cleanup()

  def test_a_headless_convert_is_run_and_the_result_is_verified_on_a_copy(self):
    """The engine verifies on a COPY by default: a LibreOffice re-save can
    alter formatting on a client-operable workbook, and the gate's
    invariant is 'verified', not 're-saved'."""
    wb = _workbook_with_status("OK", self.tmp)
    before = wb.read_bytes()
    err, verified_path = wms._recalc_workbook_via_libreoffice(
      str(wb), soffice=_fake_soffice(self.tmp), write_back=False)
    self.assertIsNone(err)
    self.assertTrue(Path(verified_path).exists())
    self.assertNotEqual(str(Path(verified_path).resolve()), str(wb.resolve()))
    self.assertEqual(wb.read_bytes(), before, "the deliverable was mutated")

  def test_write_back_replaces_the_deliverable_in_place(self):
    wb = _workbook_with_status("OK", self.tmp)
    err, verified_path = wms._recalc_workbook_via_libreoffice(
      str(wb), soffice=_fake_soffice(self.tmp), write_back=True)
    self.assertIsNone(err)
    self.assertEqual(str(Path(verified_path).resolve()), str(wb.resolve()))

  def test_a_convert_that_writes_nothing_is_an_error(self):
    wb = _workbook_with_status("OK", self.tmp)
    err, _ = wms._recalc_workbook_via_libreoffice(
      str(wb), soffice=_fake_soffice(self.tmp, behaviour="silent"))
    self.assertIsNotNone(err)
    self.assertIn("libreoffice", err)

  def test_a_crashing_convert_is_an_error(self):
    wb = _workbook_with_status("OK", self.tmp)
    err, _ = wms._recalc_workbook_via_libreoffice(
      str(wb), soffice=_fake_soffice(self.tmp, behaviour="crash"))
    self.assertIsNotNone(err)

  def test_no_soffice_anywhere_is_absent(self):
    wb = _workbook_with_status("OK", self.tmp)
    prev = os.environ.pop("LIBREOFFICE_SOFFICE", None)
    try:
      err, _ = wms._recalc_workbook_via_libreoffice(
        str(wb), soffice=None, which=lambda _n: None, candidates=())
    finally:
      if prev is not None:
        os.environ["LIBREOFFICE_SOFFICE"] = prev
    self.assertIsNotNone(err)
    self.assertEqual(wms.classify_recalc_error(err), "absent")


class ACheckThatCannotRunFails(unittest.TestCase):
  """The inverted tests asserted the opposite of this. They are gone."""

  def setUp(self):
    self._tmp = tempfile.TemporaryDirectory(prefix="wb_gate_")
    self.tmp = Path(self._tmp.name)
    self._orig_recalc = wms.recalc_workbook

  def tearDown(self):
    wms.recalc_workbook = self._orig_recalc
    self._tmp.cleanup()

  def test_no_engine_FAILS_even_when_the_cell_would_read_ok(self):
    wb = _workbook_with_status("OK", self.tmp)
    wms.recalc_workbook = lambda p, **kw: ("none", "no_recalc_engine_available: excel absent; libreoffice absent", None)
    with self.assertRaises(PostIntakePreconditionFailed) as ctx:
      wms.assert_workbook_model_status_ok(str(wb))
    self.assertEqual(ctx.exception.operation, "workbook_model_status_recalc_unavailable")

  def test_no_engine_FAILS_when_the_cell_reads_fail(self):
    """The exact case the deleted tests asserted should PASS."""
    wb = _workbook_with_status("FAIL", self.tmp)
    wms.recalc_workbook = lambda p, **kw: ("none", "no_recalc_engine_available", None)
    with self.assertRaises(PostIntakePreconditionFailed):
      wms.assert_workbook_model_status_ok(str(wb))

  def test_an_exhausted_transient_FAILS(self):
    wb = _workbook_with_status("OK", self.tmp)
    wms.recalc_workbook = lambda p, **kw: (
      "excel_com", "excel_com_failure: com_error: (-2147418111, 'Call was rejected by callee.', None, None)", None)
    with self.assertRaises(PostIntakePreconditionFailed) as ctx:
      wms.assert_workbook_model_status_ok(str(wb))
    self.assertIn("2147418111", str(ctx.exception))

  def test_a_recalculated_ok_passes(self):
    wb = _workbook_with_status("OK", self.tmp)
    wms.recalc_workbook = lambda p, **kw: ("excel_com", None, str(p))
    wms.assert_workbook_model_status_ok(str(wb))

  def test_a_recalculated_fail_raises_as_before(self):
    wb = _workbook_with_status("FAIL", self.tmp)
    wms.recalc_workbook = lambda p, **kw: ("excel_com", None, str(p))
    with self.assertRaises(PostIntakePreconditionFailed) as ctx:
      wms.assert_workbook_model_status_ok(str(wb))
    self.assertEqual(ctx.exception.operation, "workbook_model_status_fail")

  def test_verification_reads_the_engine_s_verified_path(self):
    """When an engine verifies on a copy, B2 is read from THAT file."""
    wb = _workbook_with_status("FAIL", self.tmp, name="deliverable.xlsx")
    copy_ok = _workbook_with_status("OK", self.tmp, name="verified_copy.xlsx")
    wms.recalc_workbook = lambda p, **kw: ("libreoffice", None, str(copy_ok))
    wms.assert_workbook_model_status_ok(str(wb))  # OK came from the verified copy


class TheEngineOrderIsDeterministic(unittest.TestCase):
  def setUp(self):
    self._orig_x = wms._recalc_workbook_via_excel_com
    self._orig_l = wms._recalc_workbook_via_libreoffice

  def tearDown(self):
    wms._recalc_workbook_via_excel_com = self._orig_x
    wms._recalc_workbook_via_libreoffice = self._orig_l

  def test_excel_absent_falls_through_to_libreoffice(self):
    wms._recalc_workbook_via_excel_com = lambda p, **kw: "pywin32_unavailable: ImportError"
    wms._recalc_workbook_via_libreoffice = lambda p, **kw: (None, "/tmp/verified.xlsx")
    engine, err, path = wms.recalc_workbook("x.xlsx")
    self.assertEqual((engine, err, path), ("libreoffice", None, "/tmp/verified.xlsx"))

  def test_both_absent_reports_both(self):
    wms._recalc_workbook_via_excel_com = lambda p, **kw: "pywin32_unavailable: ImportError"
    wms._recalc_workbook_via_libreoffice = lambda p, **kw: ("libreoffice_unavailable: no soffice", None)
    engine, err, path = wms.recalc_workbook("x.xlsx")
    self.assertEqual(engine, "none")
    self.assertIn("pywin32", err)
    self.assertIn("soffice", err)
    self.assertIsNone(path)

  def test_a_real_excel_failure_does_not_fall_through(self):
    """A genuine failure is a verdict about the workbook, not about the
    environment - a second engine must not be allowed to launder it."""
    wms._recalc_workbook_via_excel_com = lambda p, **kw: "excel_com_failure: com_error: (-2147352567, 'Save method of Workbook class failed', None, None)"
    called = []
    wms._recalc_workbook_via_libreoffice = lambda p, **kw: called.append(p) or (None, p)
    engine, err, _ = wms.recalc_workbook("x.xlsx")
    self.assertEqual(engine, "excel_com")
    self.assertIsNotNone(err)
    self.assertEqual(called, [])


if __name__ == "__main__":
  unittest.main()
