"""THE ARTIFACT GATE pins (Nick 2026-09-10): the audit reads the
finished docx, not the render report. Three delivered plans in a row
carried a missing figure that the completeness gate passed - because
every check interrogated the gate's own intentions. These pins prove
the audit sees the PAGE: numbering gaps, image/caption parity, and
report claims (placed or absent) that the page contradicts.
"""
from __future__ import annotations

import io
import os
import struct
import sys
import tempfile
import unittest
import zlib

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
for p in (ROOT, os.path.join(ROOT, "python")):
  if p not in sys.path:
    sys.path.insert(0, p)

import docx as _docx  # noqa: E402
from writing_phase_v2.docx_audit import audit_docx  # noqa: E402


def _tiny_png() -> io.BytesIO:
  def chunk(tag, data):
    c = struct.pack(">I", len(data)) + tag + data
    return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
  ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
  idat = zlib.compress(b"\x00\xff\x00\x00")
  raw = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
         + chunk(b"IDAT", idat) + chunk(b"IEND", b""))
  return io.BytesIO(raw)


REGISTRY = [
  {"id": "wage_positioning", "kind": "figure",
   "caption": "Stated wages against the metro wage distribution"},
  {"id": "revenue_by_line", "kind": "figure",
   "caption": "Revenue by line of business, Years 1-5"},
  {"id": "key_ratios", "kind": "table"},
]


def _build(fig_captions, n_images, table_captions=()):
  d = _docx.Document()
  for i in range(n_images):
    d.add_picture(_tiny_png())
  for cap in fig_captions:
    d.add_paragraph(cap)
  for cap in table_captions:
    d.add_table(rows=1, cols=1)
    d.add_paragraph(cap)
  f = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
  f.close()
  d.save(f.name)
  return f.name


class ArtifactAuditTests(unittest.TestCase):
  def _audit(self, path, report=None):
    try:
      return audit_docx(path, REGISTRY, report)
    finally:
      os.unlink(path)

  def test_clean_document_passes(self):
    path = _build(
      ["Figure 1 — Stated wages against the metro wage distribution",
       "Figure 2 — Revenue by line of business, Years 1-5"], 2,
      ["Table 1 — Key ratios; scheduled debt service is interest."])
    report = {"wage_positioning": {"placed": True},
              "revenue_by_line": {"placed": True},
              "key_ratios": {"placed": True}}
    self.assertEqual(self._audit(path, report), [])

  def test_skipped_figure_number_is_a_finding(self):
    path = _build(
      ["Figure 1 — Stated wages against the metro wage distribution",
       "Figure 3 — Revenue by line of business, Years 1-5"], 2)
    out = self._audit(path, {"wage_positioning": {"placed": True},
                             "revenue_by_line": {"placed": True},
                             "key_ratios": {"placed": False,
                                            "reason": "empty"}})
    self.assertTrue(any("numbering broken" in f for f in out), out)

  def test_image_without_caption_is_a_finding(self):
    path = _build(
      ["Figure 1 — Stated wages against the metro wage distribution"], 2)
    out = self._audit(path, {"wage_positioning": {"placed": True},
                             "revenue_by_line": {"placed": False,
                                                 "reason": "r"},
                             "key_ratios": {"placed": False,
                                            "reason": "empty"}})
    self.assertTrue(any("parity" in f for f in out), out)

  def test_claimed_placed_but_not_on_page(self):
    """The three-delivered-plans defect, made structural: the report
    says placed, the page disagrees, the audit says so."""
    path = _build(
      ["Figure 1 — Stated wages against the metro wage distribution"], 1)
    out = self._audit(path, {"wage_positioning": {"placed": True},
                             "revenue_by_line": {"placed": True},
                             "key_ratios": {"placed": False,
                                            "reason": "empty"}})
    self.assertTrue(any("revenue_by_line" in f and "claims PLACED" in f
                        for f in out), out)

  def test_claimed_absent_but_on_page(self):
    path = _build(
      ["Figure 1 — Stated wages against the metro wage distribution",
       "Figure 2 — Revenue by line of business, Years 1-5"], 2)
    out = self._audit(path, {"wage_positioning": {"placed": True},
                             "revenue_by_line": {"placed": False,
                                                 "reason": "single line"},
                             "key_ratios": {"placed": False,
                                            "reason": "empty"}})
    self.assertTrue(any("IS on the page" in f for f in out), out)

  def test_no_report_row_and_not_on_page(self):
    path = _build(
      ["Figure 1 — Stated wages against the metro wage distribution"], 1)
    out = self._audit(path, None)
    self.assertTrue(any("no render report row explains" in f
                        for f in out), out)


if __name__ == "__main__":
  unittest.main()
