"""THE ARTIFACT GATE (Nick 2026-09-10).

"A gate that audits its own intentions instead of the artifact is not a
gate. The artifact is the only thing a client ever sees, and it's the
only thing worth checking."

This module opens the FINISHED docx - not the render report, not the
registry's opinion of itself - and verifies:

1. Figure captions are numbered 1..F in page order with nothing skipped
   or duplicated; Table captions 1..T likewise.
2. Every image on the page has a figure caption and every figure
   caption an image (parity); every table element a table caption.
3. Every registry item the render report claims PLACED is actually on
   the page (matched by caption text), and every claimed absence is
   actually absent - a report that disagrees with the page is itself a
   finding.

Registry-table captions live in render_plan_v2.js; the snippets below
mirror them deliberately - if the renderer's caption changes without
this map, the audit fails loudly and forces the two back into sync.
"""
from __future__ import annotations

import re
from typing import Any, Dict, List, Optional, Tuple

try:
    import docx  # python-docx
except Exception:  # pragma: no cover
    docx = None

_CAP_RE = re.compile(r"^(Figure|Table)\s+(\d+)\s*[—\-–]\s*(.*)$")

# caption snippets for registry tables (render_plan_v2.js TABLES map)
_TABLE_SNIPPETS = {
    "income_statement_annual": "projected income statement",
    "balance_sheet_annual": "projected balance sheet",
    "cash_flow_annual": "projected cash flow",
    "key_ratios": "key ratios",
    "debt_schedule": "term-loan schedule",
}


def _norm(s: str) -> str:
    s = str(s or "").casefold()
    s = re.sub(r"[—–\-‐-―]", "-", s)
    s = re.sub(r"[^a-z0-9 %\-\.,]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _read_docx(path: str):
    d = docx.Document(path)
    captions: List[Tuple[str, int, str]] = []
    for p in d.paragraphs:
        m = _CAP_RE.match(p.text.strip())
        if m:
            captions.append((m.group(1), int(m.group(2)), m.group(3).strip()))
    images = len(d.inline_shapes)
    tables = len(d.tables)
    return captions, images, tables


def audit_docx(path: str, registry: List[Dict[str, Any]],
               report: Optional[Dict[str, Dict[str, Any]]] = None) -> List[str]:
    """Returns findings; empty list = the artifact matches its claims."""
    if docx is None:  # pragma: no cover
        return ["python-docx is not installed - the artifact cannot be read"]
    findings: List[str] = []
    captions, images, tables = _read_docx(path)
    fig_caps = [(n, t) for kind, n, t in captions if kind == "Figure"]
    tab_caps = [(n, t) for kind, n, t in captions if kind == "Table"]

    for label, caps in (("Figure", fig_caps), ("Table", tab_caps)):
        nums = [n for n, _ in caps]
        if nums != list(range(1, len(nums) + 1)):
            findings.append(
                "%s numbering broken: page shows %s - captions must run "
                "1..%d in order with nothing skipped"
                % (label, nums, len(nums)))

    if images != len(fig_caps):
        findings.append(
            "image/caption parity broken: %d images on the page, %d "
            "figure captions" % (images, len(fig_caps)))
    if tables != len(tab_caps):
        findings.append(
            "table/caption parity broken: %d table elements on the page, "
            "%d table captions" % (tables, len(tab_caps)))

    fig_texts = [_norm(t) for _, t in fig_caps]
    tab_texts = [_norm(t) for _, t in tab_caps]

    for it in registry:
        fid, kind = it["id"], it.get("kind")
        if kind == "figure":
            expected = _norm(it.get("caption") or fid.replace("_", " "))
            on_page = any(expected and expected in t for t in fig_texts)
        else:
            snippet = _TABLE_SNIPPETS.get(fid)
            if snippet is None:
                findings.append(
                    "%s: registry table has no caption snippet in the "
                    "artifact audit - add it (the audit cannot see this "
                    "table)" % fid)
                continue
            on_page = any(snippet in t for t in tab_texts)
        row = (report or {}).get(fid) or {}
        claimed = bool(row.get("placed")) if row else None
        if claimed is True and not on_page:
            findings.append(
                "%s: render report claims PLACED but no matching %s "
                "caption is on the page" % (fid, kind))
        elif claimed is False and on_page:
            findings.append(
                "%s: render report claims absent (%r) but the %s IS on "
                "the page" % (fid, row.get("reason"), kind))
        elif claimed is None and not on_page:
            findings.append(
                "%s: registry %s not on the page and no render report "
                "row explains it" % (fid, kind))
    return findings
