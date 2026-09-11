"""'What Changed' - the restructure's account, in the workbook itself.

Nick 2026-09-11: "A client holding a restructured workbook should be able to
see what was changed from the business they described, in the artifact
itself." The rows are post_intake_restructure.client_bounds.restructure_changes
- the same list the run email carries - so the workbook and the email can
never tell two stories.

Added ONLY to a restructured plan's workbook (export_workbook_for_row's
what_changed=), right after the Cover. Every other workbook is untouched.
"""
from __future__ import annotations

from typing import Any, Dict, List

from openpyxl.styles import Alignment

from . import design
from .cover_sheet import COVER_SHEET

WHAT_CHANGED_SHEET = "What Changed"
HEADLINE = "RESTRUCTURED PLAN - NOT THE BUSINESS AS DESCRIBED"
EXPLAINER = (
  "The plan built on the business as the client described it did not pass. The "
  "executive reshaped it, inside the client's own business, to find a version that "
  "works. None of the changes below was stated by the client. Both columns are "
  "plans built by the same model - the one from the client's description, and this "
  "restructured one; payroll includes employer costs. Every figure elsewhere in this "
  "workbook is the restructured plan's.")
COLUMNS = ("Area", "Plan as the client described it", "This restructured plan", "Note")
WIDTHS = (38, 32, 36, 60)


def build_what_changed_sheet(wb, rows: List[Dict[str, Any]]):
  index = 1 if wb.sheetnames and wb.sheetnames[0] == COVER_SHEET else 0
  ws = wb.create_sheet(WHAT_CHANGED_SHEET, index)
  for col, width in enumerate(WIDTHS, start=1):
    ws.column_dimensions[chr(64 + col)].width = width

  title = ws.cell(row=1, column=1, value=HEADLINE)
  title.font = design.font("section")
  for col in range(1, len(COLUMNS) + 1):
    ws.cell(row=1, column=col).fill = design.fill(design.NAVY)
  ws.row_dimensions[1].height = 24

  ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNS))
  note = ws.cell(row=2, column=1, value=EXPLAINER)
  note.alignment = Alignment(wrap_text=True, vertical="top")
  ws.row_dimensions[2].height = 48

  for col, label in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=4, column=col, value=label)
    cell.font = design.font("label_strong")
    cell.fill = design.fill(design.TINT_2)

  r = 5
  for item in rows or []:
    values = (item.get("area"), item.get("as_described"), item.get("restructured"), item.get("note"))
    for col, value in enumerate(values, start=1):
      cell = ws.cell(row=r, column=col, value=str(value or ""))
      cell.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
  if not rows:
    ws.cell(row=r, column=1, value="(no changes recorded)")
  ws.freeze_panes = "A5"
  return ws
